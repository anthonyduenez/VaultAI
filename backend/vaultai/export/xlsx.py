"""XLSX writer. Every text cell is stored as a string, never a formula, whatever the document said."""

from __future__ import annotations

import io
from datetime import UTC, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from vaultai.export import FORMULA_PREFIXES, TEXT_KINDS, Kind, Table

NUMBER_FORMATS: dict[Kind, str] = {
    Kind.TEXT: "@",
    Kind.ID: "@",
    Kind.MONEY: "#,##0.00",
    Kind.DATE: "yyyy-mm-dd",
    Kind.DATETIME: "yyyy-mm-dd hh:mm:ss",
    Kind.NUMBER: "0.00",
    Kind.INT: "0",
}
EXCEL_MAX_CELL_CHARS = 32767
HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="E8ECF2")


def _text(value: object) -> str:
    return ILLEGAL_CHARACTERS_RE.sub("", str(value))[:EXCEL_MAX_CELL_CHARS]


def _value(value: object, kind: Kind) -> object:
    if value is None:
        return None
    if kind in TEXT_KINDS:
        return _text(value)
    if kind is Kind.MONEY:
        return Decimal(str(value))
    if kind is Kind.DATETIME and isinstance(value, datetime):
        return value.astimezone(UTC).replace(tzinfo=None, microsecond=0)
    return value


def _write_sheet(ws: Worksheet, table: Table) -> None:
    for col_idx, column in enumerate(table.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=column.header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(col_idx)].width = column.width
    for row_idx, row in enumerate(table.rows, start=2):
        for col_idx, (column, raw) in enumerate(zip(table.columns, row, strict=True), start=1):
            value = _value(raw, column.kind)
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = value
            cell.number_format = NUMBER_FORMATS[column.kind]
            if isinstance(value, str):
                # openpyxl turns a leading "=" into a formula on assignment; force the string type back.
                cell.data_type = "s"
                if value.startswith(FORMULA_PREFIXES):
                    cell.quotePrefix = True
    ws.freeze_panes = "A2"
    if table.rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(table.columns))}{len(table.rows) + 1}"


def render_xlsx(tables: list[Table]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "VaultAI"
    for table in tables:
        _write_sheet(wb.create_sheet(table.name), table)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
