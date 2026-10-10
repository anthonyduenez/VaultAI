from __future__ import annotations

import csv
import hashlib
import io
import zipfile
from datetime import UTC, datetime

import openpyxl
from fastapi.testclient import TestClient

from tests.surfaces.conftest import FakeCore, field_row, insert_document
from vaultai.context import AppContext

APPROVED_AT = datetime(2026, 10, 9, 14, 30, tzinfo=UTC)
INJECTIONS = ['=HYPERLINK("http://x","y")', "+1+1", "-2+3", "@SUM(A1)"]


def approved(ctx: AppContext, **kwargs: object) -> str:
    base = {"status": "approved", "approval_kind": "human", "approved_by": "sam", "approved_at": APPROVED_AT}
    return insert_document(ctx, **{**base, **kwargs})  # type: ignore[arg-type]


def export(client: TestClient, **filters: object) -> tuple[dict, bytes]:
    job = client.post("/api/exports", json=filters).json()
    download = client.get(f"/api/exports/{job['id']}/download")
    assert download.status_code == 200
    return job, download.content


def workbook(data: bytes) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(data))


def rows(ws) -> list[dict[str, object]]:  # type: ignore[no-untyped-def]
    header = [c.value for c in ws[1]]
    return [dict(zip(header, (c.value for c in row), strict=True)) for row in ws.iter_rows(min_row=2)]


def test_default_export_contains_only_approved_records_with_typed_cells(
    client: TestClient, app_ctx: AppContext, fake_core: FakeCore
) -> None:
    check = approved(
        app_ctx, filename="check_4521.pdf", doc_type="check", counterparty="Northwind Traders", doc_date="2026-03-04",
        total="5328.00", primary_id="000123456789",
        fields=[field_row("check_number", "000123456789", value_type="identifier", label="Check number"),
                field_row("amount", "5328.00", value_type="money", label="Amount", origin="model")],
    )
    pending = insert_document(app_ctx, filename="pending.pdf", status="needs_review", counterparty="Pending Co")
    client.patch(f"/api/documents/{check}/fields/amount", json={"value": "5328.00"})
    client.post(f"/api/documents/{check}/review", json={"action": "approve"})

    job, data = export(client)
    wb = workbook(data)

    assert wb.sheetnames == ["Export Info", "Approved Records", "Source Evidence", "Processing Metadata",
                             "Validation Warnings", "Obligations and Deadlines"]
    records = rows(wb["Approved Records"])
    assert [r["Document ID"] for r in records] == [check]
    record = records[0]
    assert (record["File"], record["Type"], record["Counterparty"], record["Identifier"], record["Total"],
            record["Currency"], record["Approval"], record["Document Date"], record["Human Edits"]) == (
        "check_4521.pdf", "Check", "Northwind Traders", "000123456789", 5328, "USD", "Human",
        datetime(2026, 3, 4), 1,
    )
    sheet = wb["Approved Records"]
    identifier = sheet.cell(row=2, column=7)
    total = sheet.cell(row=2, column=8)
    assert (identifier.data_type, identifier.number_format) == ("s", "@")
    assert (total.data_type, total.number_format) == ("n", "#,##0.00")
    assert sheet.cell(row=2, column=6).number_format == "yyyy-mm-dd"
    assert sheet.freeze_panes == "A2"

    evidence = {r["Field"]: r for r in rows(wb["Source Evidence"])}
    assert (evidence["Check number"]["Value Origin"], evidence["Amount"]["Value Origin"]) == ("AI-rules", "Human-edited")
    assert (evidence["Amount"]["AI Value"], evidence["Amount"]["Edited By"], evidence["Amount"]["Line ID"]) == (
        "5328.00", "local-user", "p1.l1",
    )
    info = {r["Field"]: r["Value"] for r in rows(wb["Export Info"])}
    assert (info["Export ID"], info["Approved Records"], info["Exported By"]) == (job["id"], "1", "local-user")
    assert info["Notice"].endswith("VaultAI is a document assistant, not legal or financial advice.")

    assert job["record_count"] == 1
    assert job["sha256"] == hashlib.sha256(data).hexdigest()
    assert client.get(f"/api/documents/{check}").json()["status"] == "exported"
    assert client.get(f"/api/documents/{pending}").json()["status"] == "needs_review"
    history = client.get(f"/api/documents/{check}").json()["history"]
    assert (history[0]["action"], history[0]["after"], history[0]["note"]) == ("export", "exported", f"Export {job['id']}")


def test_needs_review_rows_go_to_their_own_sheet(client: TestClient, app_ctx: AppContext) -> None:
    ok = approved(app_ctx, filename="ok.pdf")
    pending = insert_document(app_ctx, filename="pending.pdf", status="ready")

    preview = client.post("/api/exports/preview", json={"include_needs_review": True}).json()
    assert (preview["record_count"], preview["needs_review_count"]) == (1, 1)
    assert preview["warnings"] == [
        "1 record that still needs review goes to a separate Needs Review sheet and not into Approved Records."
    ]
    assert client.post("/api/exports/preview", json={}).json()["warnings"] == [
        "1 matching record still needs review and is not included."
    ]

    _, data = export(client, include_needs_review=True, include_evidence=False)
    wb = workbook(data)
    assert wb.sheetnames == ["Export Info", "Approved Records", "Needs Review", "Processing Metadata",
                             "Validation Warnings", "Obligations and Deadlines"]
    assert [r["Document ID"] for r in rows(wb["Approved Records"])] == [ok]
    assert [(r["Document ID"], r["Approval"]) for r in rows(wb["Needs Review"])] == [(pending, None)]
    assert client.get(f"/api/documents/{pending}").json()["status"] == "ready"


def test_filters_by_type_and_date(client: TestClient, app_ctx: AppContext) -> None:
    march = approved(app_ctx, doc_type="invoice", doc_date="2026-03-10")
    approved(app_ctx, doc_type="invoice", doc_date="2026-05-10")
    approved(app_ctx, doc_type="receipt", doc_date="2026-03-11")
    approved(app_ctx, doc_type="invoice", doc_date=None)

    preview = client.post(
        "/api/exports/preview", json={"doc_types": ["invoice"], "date_from": "2026-03-01", "date_to": "2026-03-31"}
    ).json()
    assert [r["doc_id"] for r in preview["rows"]] == [march]
    assert preview["by_type"] == [{"doc_type": "invoice", "count": 1}]


def test_formula_payloads_are_stored_as_strings(client: TestClient, app_ctx: AppContext) -> None:
    approved(
        app_ctx, counterparty=INJECTIONS[0], primary_id=INJECTIONS[1],
        fields=[field_row(f"f{i}", payload) for i, payload in enumerate(INJECTIONS)],
        issues=[("risky_clause", "warning", INJECTIONS[3])],
    )
    _, data = export(client)

    archive = zipfile.ZipFile(io.BytesIO(data))
    xml = "".join(archive.read(n).decode() for n in archive.namelist() if n.startswith("xl/"))
    assert "<f>" not in xml and "<f " not in xml
    for payload in ["=HYPERLINK(", "+1+1", "-2+3", "@SUM(A1)"]:
        assert payload in xml
    record = rows(workbook(data)["Approved Records"])[0]
    assert (record["Counterparty"], record["Identifier"]) == ('=HYPERLINK("http://x","y")', "+1+1")


def test_csv_export_defuses_formulas_and_starts_with_bom(client: TestClient, app_ctx: AppContext) -> None:
    doc_id = approved(app_ctx, filename="@evil.pdf", counterparty="-2+3", primary_id="000123", total="12.50",
                      doc_date="2026-01-02")
    job, data = export(client, format="csv")

    assert job["filename"].endswith(".csv")
    assert data.startswith(b"\xef\xbb\xbf")
    parsed = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
    assert [(r["Document ID"], r["File"], r["Counterparty"], r["Identifier"], r["Total"], r["Document Date"])
            for r in parsed] == [(doc_id, "'@evil.pdf", "'-2+3", "000123", "12.50", "2026-01-02")]


def test_export_list_newest_first(client: TestClient, app_ctx: AppContext) -> None:
    approved(app_ctx)
    first, _ = export(client)
    second, _ = export(client, format="csv")
    listed = client.get("/api/exports").json()
    assert [j["id"] for j in listed] == [second["id"], first["id"]]
    assert (listed[0]["record_count"], listed[1]["record_count"]) == (1, 1)
